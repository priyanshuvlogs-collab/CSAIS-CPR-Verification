"""CSV / XLSX / PDF report generation."""
import csv
import io
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from app.services.expiry import EXPIRED, EXPIRING, UNKNOWN, VALID

COLUMNS = [
    ("id", "ID"),
    ("guard_name", "Guard name"),
    ("guard_id", "Guard ID"),
    ("site", "Site"),
    ("holder_name", "Name on certificate"),
    ("issuer", "Issuer"),
    ("cert_id", "Certificate #"),
    ("issue_date", "Issue date"),
    ("expiry_date", "Expiry date"),
    ("expiry_inferred", "Expiry inferred"),
    ("status", "Expiry status"),
    ("days_left", "Days left"),
    ("tamper_level", "Tamper risk"),
    ("tamper_score", "Tamper score"),
    ("tamper_summary", "Tamper indicators"),
    ("reviewed", "Reviewed"),
    ("uploaded_at", "Uploaded"),
]

STATUS_FILL = {EXPIRED: "F8D7DA", EXPIRING: "FFF3CD", VALID: "D1E7DD", UNKNOWN: "E2E3E5"}
RISK_FILL = {"High": "F8D7DA", "Medium": "FFF3CD", "Low": "D1E7DD"}


def build_rows(items) -> list[dict]:
    """items: iterable of (submission, ExpiryInfo)."""
    rows = []
    for sub, info in items:
        rows.append({
            "id": sub.id,
            "guard_name": sub.guard_name,
            "guard_id": sub.guard_id,
            "site": sub.site or "",
            "holder_name": sub.holder_name or "",
            "issuer": sub.issuer or "",
            "cert_id": sub.cert_id or "",
            "issue_date": sub.issue_date.isoformat() if sub.issue_date else "",
            "expiry_date": sub.expiry_date.isoformat() if sub.expiry_date else "",
            "expiry_inferred": "Yes" if sub.expiry_inferred else "",
            "status": info.status,
            "days_left": info.days_left if info.days_left is not None else "",
            "tamper_level": sub.tamper_level or "",
            "tamper_score": sub.tamper_score if sub.tamper_score is not None else "",
            "tamper_summary": "; ".join(r["detail"] for r in (sub.tamper_reasons or [])),
            "reviewed": "Yes" if sub.reviewed else "",
            "uploaded_at": sub.uploaded_at.strftime("%Y-%m-%d %H:%M") if sub.uploaded_at else "",
        })
    return rows


def summary(rows: list[dict]) -> dict:
    s = {EXPIRED: 0, EXPIRING: 0, VALID: 0, UNKNOWN: 0, "High": 0, "Medium": 0, "total": len(rows)}
    for r in rows:
        s[r["status"]] = s.get(r["status"], 0) + 1
        if r["tamper_level"] in ("High", "Medium"):
            s[r["tamper_level"]] += 1
    return s


def to_csv(rows: list[dict]) -> bytes:
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow([label for _, label in COLUMNS])
    for r in rows:
        w.writerow([r[key] for key, _ in COLUMNS])
    return buf.getvalue().encode("utf-8-sig")  # BOM so Excel opens UTF-8 correctly


def to_xlsx(rows: list[dict]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "CPR Certificates"
    ws.append([label for _, label in COLUMNS])
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F3A5F")
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    keys = [k for k, _ in COLUMNS]
    status_col = keys.index("status") + 1
    risk_col = keys.index("tamper_level") + 1
    for r in rows:
        ws.append([r[k] for k in keys])
        row_idx = ws.max_row
        if fill := STATUS_FILL.get(r["status"]):
            ws.cell(row_idx, status_col).fill = PatternFill("solid", fgColor=fill)
        if fill := RISK_FILL.get(r["tamper_level"]):
            ws.cell(row_idx, risk_col).fill = PatternFill("solid", fgColor=fill)
    widths = {"tamper_summary": 60, "guard_name": 22, "holder_name": 22, "issuer": 24}
    for i, k in enumerate(keys, start=1):
        ws.column_dimensions[get_column_letter(i)].width = widths.get(k, 14)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    s = summary(rows)
    ws2 = wb.create_sheet("Summary")
    ws2.append(["Generated", date.today().isoformat()])
    for label, key in [("Total", "total"), ("Expired", EXPIRED), ("Expiring within 6 months", EXPIRING),
                       ("Valid", VALID), ("Unknown expiry", UNKNOWN), ("High tamper risk", "High"),
                       ("Medium tamper risk", "Medium")]:
        ws2.append([label, s[key]])
    ws2.column_dimensions["A"].width = 28

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def to_pdf(rows: list[dict], warning_months: int = 6, filters: dict | None = None) -> bytes:
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4), leftMargin=12 * mm, rightMargin=12 * mm,
                            topMargin=12 * mm, bottomMargin=12 * mm, title="CPR Certificate Report")
    styles = getSampleStyleSheet()
    small = styles["BodyText"].clone("small", fontSize=7.5, leading=9)
    story = [Paragraph("CPR Certificate Verification Report", styles["Title"]),
             Paragraph(f"Generated {date.today():%d %B %Y}", styles["Normal"])]
    active = {k: v for k, v in (filters or {}).items() if v}
    if active:
        story.append(Paragraph("Filtered by: " + ", ".join(f"{k} = {v}" for k, v in active.items()), styles["Normal"]))
    story.append(Spacer(1, 6 * mm))

    s = summary(rows)
    summary_data = [["Total", "Expired", f"Expiring ≤{warning_months} mo", "Valid", "Unknown", "High risk", "Medium risk"],
                    [s["total"], s[EXPIRED], s[EXPIRING], s[VALID], s[UNKNOWN], s["High"], s["Medium"]]]
    t = Table(summary_data, hAlign="LEFT")
    t.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F3A5F")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTSIZE", (0, 1), (-1, 1), 14),
        ("ALIGN", (0, 0), (-1, -1), "CENTER"),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
        ("BOTTOMPADDING", (0, 1), (-1, 1), 6),
    ]))
    story += [t, Spacer(1, 6 * mm)]

    def section(title, subset, cols):
        story.append(Paragraph(f"{title} ({len(subset)})", styles["Heading2"]))
        if not subset:
            story.append(Paragraph("None.", styles["Normal"]))
            story.append(Spacer(1, 4 * mm))
            return
        head = [label for _, label in cols]
        data = [head] + [[Paragraph(str(r[k]), small) for k, _ in cols] for r in subset]
        tbl = Table(data, repeatRows=1)
        tbl.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F3A5F")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTSIZE", (0, 0), (-1, 0), 8),
            ("GRID", (0, 0), (-1, -1), 0.25, colors.grey),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F4F6F9")]),
        ]))
        story.append(tbl)
        story.append(Spacer(1, 5 * mm))

    expiry_cols = [("guard_name", "Guard"), ("guard_id", "Guard ID"), ("site", "Site"), ("issuer", "Issuer"),
                   ("cert_id", "Cert #"), ("expiry_date", "Expiry"), ("days_left", "Days left"),
                   ("tamper_level", "Tamper risk")]
    by_days = lambda r: r["days_left"] if r["days_left"] != "" else 0  # noqa: E731
    section("Expired", sorted([r for r in rows if r["status"] == EXPIRED], key=by_days), expiry_cols)
    section(f"Expiring within {warning_months} months",
            sorted([r for r in rows if r["status"] == EXPIRING], key=by_days), expiry_cols)
    section("Unknown expiry – needs manual review", [r for r in rows if r["status"] == UNKNOWN],
            [("guard_name", "Guard"), ("guard_id", "Guard ID"), ("site", "Site"), ("uploaded_at", "Uploaded")])

    flagged = sorted([r for r in rows if r["tamper_level"] in ("High", "Medium")],
                     key=lambda r: -(r["tamper_score"] or 0))
    if flagged:
        story.append(PageBreak())
    section("Tamper risk indicators (High / Medium)", flagged,
            [("guard_name", "Guard"), ("guard_id", "Guard ID"), ("tamper_level", "Risk"),
             ("tamper_score", "Score"), ("tamper_summary", "Indicators")])
    story.append(Paragraph(
        "Tamper indicators are automated signals for review, not proof of forgery. Confirm flagged certificates "
        "directly with the issuing organisation (e.g. AHA eCard or Red Cross certificate lookup).", small))

    doc.build(story)
    return buf.getvalue()
